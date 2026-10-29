"""sitemap.xlsx: Pages, Broken links, All links, Skipped, Summary."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import __version__
from ..linkstatus import OK, classify
from ..model import CrawlResult

HEAD_FILL = PatternFill("solid", fgColor="1F2937")
HEAD_FONT = Font(bold=True, color="FFFFFF")
DECISION_FILL = PatternFill("solid", fgColor="FEF3C7")
LINK_FONT = Font(color="0563C1", underline="single")


def _header(ws, headers: list[str], widths: list[int]) -> None:
    ws.append(headers)
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def write_xlsx(result: CrawlResult, path: Path) -> dict:
    wb = Workbook()

    # ---- Pages
    ws = wb.active
    ws.title = "Pages"
    headers = ["Index", "Title", "URL", "Parent page", "Navigation path", "Depth", "Word count",
               "Desktop screenshot file", "Mobile screenshot file", "Keep", "Rewrite", "Eliminate"]
    _header(ws, headers, [7, 40, 50, 36, 45, 7, 11, 36, 36, 9, 9, 10])
    by_url = {p.url: p for p in result.pages}
    by_url.update({p.final_url: p for p in result.pages})
    for p in result.pages:
        if p.parent and p.parent in by_url:
            parent = by_url[p.parent].short_title or by_url[p.parent].title
        elif p.found_via not in ("link", "start"):
            parent = f"(found via {p.found_via}, not linked from crawled pages)"
        else:
            parent = ""
        ws.append([p.index, p.title, p.final_url, parent, p.nav_path, p.depth, p.word_count,
                   p.desktop_png, p.mobile_png, "", "", ""])
        r = ws.max_row
        ws.cell(row=r, column=3).hyperlink = p.final_url
        ws.cell(row=r, column=3).font = LINK_FONT
        for col, rel in ((8, p.desktop_png), (9, p.mobile_png)):
            if rel:
                ws.cell(row=r, column=col).hyperlink = rel
                ws.cell(row=r, column=col).font = LINK_FONT
        for col in (10, 11, 12):
            ws.cell(row=r, column=col).fill = DECISION_FILL

    # ---- Broken links
    wb_ws = wb.create_sheet("Broken links")
    _header(wb_ws, ["Source page", "Link URL", "Link text", "HTTP status or error", "Internal/External", "Assessment", "Note",
                    "Pages linking to it"], [50, 60, 30, 20, 10, 16, 60, 10])
    problems = [(l, *classify(l)) for l in result.links]
    problems = [(l, cat, note) for l, cat, note in problems if cat != OK]
    # sort: confirmed broken first, internal before external
    problems.sort(key=lambda x: (x[1] != "verify manually" and 0 or 1, x[1] != "broken", not x[0].internal, x[0].url))
    broken_urls = set()
    verify_urls = set()
    seen_verify: dict[str, int] = {}
    for l, cat, note in problems:
        if cat == "verify manually":
            # one row per URL (a social link repeated on every page would otherwise fill the sheet)
            if l.url in seen_verify:
                wb_ws.cell(row=seen_verify[l.url], column=8).value += 1
                continue
            wb_ws.append([l.source, l.url, l.text, l.status, "internal" if l.internal else "external", cat, note, 1])
            seen_verify[l.url] = wb_ws.max_row
            verify_urls.add(l.url)
            continue
        wb_ws.append([l.source, l.url, l.text, l.status, "internal" if l.internal else "external", cat, note, 1])
        broken_urls.add(l.url)
    # pages that failed during the crawl itself (timeouts, HTTP errors on queued pages)
    crawl_failures = 0
    for s in result.skipped:
        if s["reason"].startswith(("HTTP", "timeout", "error")) and s["url"] not in broken_urls and s["url"] not in verify_urls:
            wb_ws.append(["(crawler)", s["url"], "", s["reason"], "internal", "broken", "page could not be loaded during the crawl", 1])
            broken_urls.add(s["url"])
            crawl_failures += 1
    broken = [x for x in problems if x[1] != "verify manually"] + [None] * crawl_failures
    verify = [x for x in problems if x[1] == "verify manually"]

    # ---- All links
    al = wb.create_sheet("All links")
    _header(al, ["Source page", "Link URL", "Link text", "Internal/External", "HTTP status"], [50, 60, 30, 10, 14])
    for l in result.links:
        al.append([l.source, l.url, l.text, "internal" if l.internal else "external", l.status])

    # ---- Skipped
    if result.skipped:
        sk = wb.create_sheet("Skipped URLs")
        _header(sk, ["URL", "Reason"], [70, 50])
        for s in result.skipped:
            sk.append([s["url"], s["reason"]])

    # ---- Summary
    checked = sum(1 for l in result.links if l.status and l.status != "not checked")
    distinct_links = len({l.url for l in result.links})
    stats = {
        "pages": len(result.pages),
        "broken_links": len(broken),
        "broken_urls": len(broken_urls),
        "verify_links": len(verify),
        "verify_urls": len(verify_urls),
        "links_seen": len(result.links),
        "distinct_links": distinct_links,
        "links_checked": checked,
        "skipped": len(result.skipped),
    }
    sm = wb.create_sheet("Summary")
    sm.column_dimensions["A"].width = 28
    sm.column_dimensions["B"].width = 70
    rows = [
        ("Site", result.options.get("site_name") or result.origin),
        ("Start URL", result.start_url),
        ("Crawl date", result.crawled_at),
        ("Page count", stats["pages"]),
        ("Broken-link count", stats["broken_links"]),
        ("Distinct broken URLs", stats["broken_urls"]),
        ("Links to verify manually", f"{stats['verify_links']} ({stats['verify_urls']} distinct URLs; external sites that block automated checks, timeouts)"),
        ("Links seen (page x target)", stats["links_seen"]),
        ("Distinct link targets", distinct_links),
        ("Skipped URLs", stats["skipped"]),
        ("Sample mode", f"yes, first {result.sample} pages" if result.sample else "no, full crawl"),
        ("Mobile screenshots", "yes" if result.options.get("mobile", True) else "no"),
        ("External links checked", "yes" if result.options.get("check_external", True) else "no"),
        ("Request delay", f"{result.options.get('delay_ms', 1000)} ms"),
        ("Tool", f"site-archive {__version__}"),
    ]
    for k, v in rows:
        sm.append([k, v])
        sm.cell(row=sm.max_row, column=1).font = Font(bold=True)
    sm.cell(row=2, column=2).hyperlink = result.start_url
    sm.cell(row=2, column=2).font = LINK_FONT

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return stats
