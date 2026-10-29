"""content.docx: all visible page text, one section per page, fully editable."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from ..model import Block, CrawlResult

GREY = RGBColor(0x6B, 0x72, 0x80)
LINK_BLUE = RGBColor(0x05, 0x63, 0xC1)
LABELS = {"link": "Link", "cta": "Call to action", "button": "Button"}
REGION_NAMES = {"nav": "Navigation", "header": "Header", "footer": "Footer", "": "Main content"}


def _add_hyperlink(paragraph, text: str, url: str | None = None, anchor: str | None = None):
    hyperlink = OxmlElement("w:hyperlink")
    if anchor:
        hyperlink.set(qn("w:anchor"), anchor)
    else:
        r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_bookmark(paragraph, name: str, bid: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    paragraph._p.append(start)
    paragraph._p.append(end)


def _ensure_styles(doc) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    for lvl in range(1, 5):
        name = f"Site H{lvl}"
        if name not in [s.name for s in styles]:
            st = styles.add_style(name, 1)  # WD_STYLE_TYPE.PARAGRAPH
            st.base_style = styles[f"Heading {lvl + 1}"]
            st.quick_style = True
    if "Site Label" not in [s.name for s in styles]:
        st = styles.add_style("Site Label", 1)
        st.base_style = styles["Normal"]
        st.font.size = Pt(9)
        st.font.color.rgb = GREY
        st.font.italic = True


def _write_block(doc, b: Block) -> None:
    t = b.type
    if t in ("h1", "h2", "h3"):
        doc.add_paragraph(b.text, style=f"Site H{t[1]}")
    elif t in ("h4", "h5", "h6"):
        doc.add_paragraph(b.text, style="Site H4")
    elif t == "li":
        doc.add_paragraph(b.text, style="List Bullet")
    elif t in LABELS:
        p = doc.add_paragraph()
        lab = p.add_run(f"{LABELS[t]}: ")
        lab.font.color.rgb = GREY
        lab.font.size = Pt(9)
        lab.italic = True
        r = p.add_run(b.text)
        if t == "cta":
            r.bold = True
    else:
        for line in b.text.split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip())


def write_docx(result: CrawlResult, path: Path, shared_blocks: list[Block] | None = None) -> None:
    doc = Document()
    _ensure_styles(doc)
    site = result.options.get("site_name") or result.origin.split("//", 1)[-1]
    pages = result.pages

    doc.add_heading(f"Website content archive: {site}", 0)
    meta = doc.add_paragraph()
    meta.add_run("Source: ").bold = True
    _add_hyperlink(meta, result.start_url, url=result.start_url)
    meta.add_run(f"\nCrawled: {result.crawled_at}\nPages archived: {len(pages)}")
    if result.sample:
        meta.add_run(f"\nSAMPLE: first {result.sample} pages only").bold = True
    doc.add_paragraph(
        "Every page below lists the text visible to a visitor, in reading order. Headings keep their "
        "levels (Site H1, Site H2, Site H3 styles). Links, buttons and calls to action are labelled. "
        "Text is editable; nothing here is an image.", style="Site Label")

    doc.add_heading("Contents", 1)
    for p in pages:
        para = doc.add_paragraph(style="List Number")
        _add_hyperlink(para, f"{p.index}. {p.short_title or p.title}", anchor=f"page_{p.index}")
        para.add_run(f"  ({p.final_url})").font.color.rgb = GREY

    for p in pages:
        doc.add_page_break()
        h = doc.add_heading(f"{p.index}. {p.title}", 1)
        _add_bookmark(h, f"page_{p.index}", 100 + p.index)
        u = doc.add_paragraph()
        u.add_run("URL: ").bold = True
        _add_hyperlink(u, p.final_url, url=p.final_url)
        info = doc.add_paragraph(style="Site Label")
        info.add_run(f"Navigation path: {p.nav_path}   |   Depth: {p.depth}   |   Words: {p.word_count}")
        if p.desktop_png:
            info.add_run(f"\nScreenshots: {p.desktop_png}" + (f", {p.mobile_png}" if p.mobile_png else ""))

        has_regions = any(b.region for b in p.blocks)
        current_region = None
        for b in p.blocks:
            if has_regions and b.region != current_region:
                current_region = b.region
                doc.add_paragraph(f"[{REGION_NAMES.get(b.region, b.region)}]", style="Site Label")
            _write_block(doc, b)
        if not p.blocks:
            doc.add_paragraph("(no visible text extracted)", style="Site Label")

    if shared_blocks:
        doc.add_page_break()
        doc.add_heading("Shared elements (navigation, header, footer)", 1)
        doc.add_paragraph(
            "These texts appear on almost every page and were removed from the individual page sections "
            "above (--dedupe-boilerplate). They are listed once here.", style="Site Label")
        for b in shared_blocks:
            _write_block(doc, b)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
